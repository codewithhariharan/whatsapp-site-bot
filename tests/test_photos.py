"""Site photos: storing them as they arrive, and /ask returning them in Excel."""
import asyncio
import base64
import io

import pytest
from openpyxl import load_workbook
from PIL import Image

import commands as cmd
import database as db
import excel_generator as xls
import ingest_batch as ib
import message_handler as mh
import photos
from config import settings

SITE_GROUP = "120363021760406818@g.us"
TUNNEL_GROUP = "120363111111111111@g.us"
MASTER_GROUP = "120363999999999999@g.us"


def jpeg(width=3000, height=2000, color=(120, 90, 60), orientation=None, mode="RGB"):
    img = Image.new(mode, (width, height), color if mode == "RGB" else color + (255,))
    out = io.BytesIO()
    if orientation:
        exif = Image.Exif()
        exif[0x0112] = orientation          # EXIF Orientation
        img.save(out, format="JPEG", exif=exif.tobytes())
    elif mode == "RGBA":
        img.save(out, format="PNG")
    else:
        img.save(out, format="JPEG")
    return out.getvalue()


# ── Preparing an image ────────────────────────────────────────────────────────

class TestPrepare:
    def test_a_large_photo_is_shrunk_to_the_long_edge(self):
        data, w, h = photos.prepare_image(jpeg(4000, 3000))
        assert (w, h) == (1600, 1200)
        assert Image.open(io.BytesIO(data)).format == "JPEG"

    def test_a_small_photo_is_not_enlarged(self):
        _, w, h = photos.prepare_image(jpeg(800, 600))
        assert (w, h) == (800, 600)

    def test_a_portrait_shot_is_stored_upright(self):
        # Phones save a portrait photo as landscape pixels plus "rotate 90".
        _, w, h = photos.prepare_image(jpeg(3000, 2000, orientation=6))
        assert h > w

    def test_a_png_with_transparency_becomes_a_jpeg(self):
        data, _, _ = photos.prepare_image(jpeg(400, 300, mode="RGBA"))
        assert Image.open(io.BytesIO(data)).mode == "RGB"


@pytest.mark.parametrize("question, expected", [
    ("show me the picture of Exit 3 in the last two months", True),
    ("any photos of P46 this week?", True),
    ("pics of CCW2", True),
    ("images from yesterday", True),
    ("when was U3-38 cast?", False),
    ("what happened at Exit 3?", False),
])
def test_wants_photos(question, expected):
    assert photos.wants_photos(question) is expected


# ── Arrival ───────────────────────────────────────────────────────────────────

@pytest.fixture
def arrival(monkeypatch):
    w = {"stored": [], "queued": []}

    def insert(**kw):
        w["stored"].append(kw)
        return True

    monkeypatch.setattr(db, "insert_site_photo", insert)
    monkeypatch.setattr(db, "enqueue_message", lambda *a: w["queued"].append(a))
    monkeypatch.setattr(settings, "TUNNEL_GROUP_IDS", TUNNEL_GROUP)
    monkeypatch.setattr(settings, "MASTER_GROUP_IDS", MASTER_GROUP)

    async def quiet(*_a, **_k):
        pass

    import tunnel_handler as th
    monkeypatch.setattr(th, "send_message", quiet)
    return w


def arrive(group_id, text, image=None, message_id="WA1"):
    asyncio.run(mh.handle_message(group_id, "Ali", "6591234567", text,
                                  image=image, message_id=message_id))


class TestArrival:
    def test_a_captioned_site_photo_is_stored_and_its_caption_queued(self, arrival):
        arrive(SITE_GROUP, "Exit 3: waterproofing works", jpeg())
        photo = arrival["stored"][0]
        assert photo["group_id"] == SITE_GROUP
        assert photo["caption"] == "Exit 3: waterproofing works"
        assert photo["wa_message_id"] == "WA1"
        assert photo["width"] == 1600
        assert arrival["queued"] == [(SITE_GROUP, "Ali", "6591234567",
                                      "Exit 3: waterproofing works")]

    def test_an_album_photo_without_caption_is_stored_and_nothing_is_queued(self, arrival):
        arrive(SITE_GROUP, "", jpeg())
        assert len(arrival["stored"]) == 1
        assert arrival["stored"][0]["caption"] is None
        assert arrival["queued"] == []

    @pytest.mark.parametrize("group", [TUNNEL_GROUP, MASTER_GROUP])
    def test_tunnel_and_master_groups_keep_no_photos(self, arrival, group):
        arrive(group, "", jpeg())
        assert arrival["stored"] == []

    def test_a_photo_that_will_not_store_does_not_cost_the_caption_its_log(
            self, arrival, monkeypatch):
        def broken(**_kw):
            raise RuntimeError("db down")

        monkeypatch.setattr(db, "insert_site_photo", broken)
        arrive(SITE_GROUP, "Exit 3: waterproofing works", jpeg())
        assert len(arrival["queued"]) == 1

    def test_the_ingest_endpoint_accepts_a_photo_with_no_caption(self, monkeypatch):
        from fastapi.testclient import TestClient

        import main
        seen = []

        async def handle(group_id, sender_name, sender_number, text, image=None,
                         message_id=None):
            seen.append((text, len(image or b""), message_id))

        monkeypatch.setattr(main, "handle_message", handle)
        client = TestClient(main.app)
        headers = {"X-Bridge-Secret": settings.BRIDGE_SHARED_SECRET}
        body = {"group_id": SITE_GROUP, "sender_number": "65", "text": "",
                "image_base64": base64.b64encode(b"abc").decode(), "message_id": "WA9"}
        assert client.post("/baileys/incoming", json=body, headers=headers).json() == {"status": "ok"}
        assert seen == [("", 3, "WA9")]
        empty = {"group_id": SITE_GROUP, "text": ""}
        assert client.post("/baileys/incoming", json=empty, headers=headers).json() == {"status": "ignored"}


# ── Linking runs with every batch ─────────────────────────────────────────────

def test_each_batch_links_photos_to_the_logs_it_filed(monkeypatch):
    linked = []
    monkeypatch.setattr(db, "get_pending_messages", lambda cutoff=None, group_ids=None: [
        {"id": 1, "group_id": SITE_GROUP, "sender_name": "Ali", "sender_number": "65",
         "text": "Exit 3: waterproofing", "received_at": "2026-09-30T10:00:00+08:00"}])
    monkeypatch.setattr(db, "mark_message", lambda *a, **k: None)
    monkeypatch.setattr(db, "upsert_group", lambda *a, **k: None)
    monkeypatch.setattr(db, "insert_log", lambda **k: None)
    monkeypatch.setattr(db, "link_site_photos", lambda ids: linked.append(ids) or 1)
    monkeypatch.setattr(ib, "classify_and_parse",
                        lambda t: {"type": "log", "data": {"main_location": "Exit 3"}})
    monkeypatch.setattr(settings, "TUNNEL_GROUP_IDS", "")
    asyncio.run(ib.run_batch(group_ids=[SITE_GROUP]))
    assert linked == [[SITE_GROUP]]


# ── Finding ───────────────────────────────────────────────────────────────────

PLAN = ('{"sql": "SELECT 1", "locations": ["Exit 3", "U3-12", "RS"], '
        '"period": "30 Jul - 30 Sep 2026"}')


def test_find_photos_groups_by_location_in_the_order_asked(monkeypatch):
    seen = {}
    monkeypatch.setattr(photos.llm, "generate", lambda *a, **k: PLAN)

    def run(sql, group_id, **kw):
        seen.update(kw, group_id=group_id)
        return [{"tab": "Exit 3", "photo_id": 7}, {"tab": "Exit 3", "photo_id": 9},
                {"tab": "U3-12", "photo_id": 9}, {"tab": "U3-12", "photo_id": 12}], False

    monkeypatch.setattr(db, "run_readonly_query", run)
    monkeypatch.setattr(db, "get_site_photos", lambda g, ids: [
        {"id": i, "image": b"x"} for i in sorted(ids)])

    sections, period, capped = photos.find_photos(SITE_GROUP, "pictures of Exit 3, U3-12 and RS")
    assert seen["group_id"] == SITE_GROUP and seen["scope_to_group"] is True
    assert seen["tables"] == ("daily_logs", "site_photos")
    assert [(loc, [p["id"] for p in found]) for loc, found in sections] == [
        ("Exit 3", [7, 9]),
        ("U3-12", [9, 12]),       # a photo matching both locations is on both tabs
        ("RS", []),               # asked about, none found: still a tab
    ]
    assert period == "30 Jul - 30 Sep 2026"
    assert capped is False


def test_the_caps_apply_per_location_and_in_all(monkeypatch):
    monkeypatch.setattr(photos, "MAX_PER_LOCATION", 2)
    monkeypatch.setattr(photos, "MAX_PHOTOS", 3)
    monkeypatch.setattr(photos.llm, "generate", lambda *a, **k: PLAN)
    monkeypatch.setattr(db, "run_readonly_query", lambda *a, **k: (
        [{"tab": "Exit 3", "photo_id": i} for i in (1, 2, 3)]
        + [{"tab": "U3-12", "photo_id": i} for i in (4, 5)], False))
    monkeypatch.setattr(db, "get_site_photos", lambda g, ids: [{"id": i} for i in ids])
    sections, _, capped = photos.find_photos(SITE_GROUP, "q")
    assert [[p["id"] for p in found] for _loc, found in sections] == [[1, 2], [4], []]
    assert capped is True


# ── The Excel sheet ───────────────────────────────────────────────────────────

def _row(color, sent, loc, orientation=(1600, 1200), **extra):
    w, h = orientation
    return {"id": hash((color, sent)), "image": photos.prepare_image(jpeg(w, h, color))[0],
            "sent_at": sent, "main_location": loc, "sender_name": "Ali", **extra}


def test_one_tab_per_location_one_photo_per_row():
    sections = [
        ("Exit 3", [_row((1, 2, 3), "2026-08-01T09:15:00+08:00", "Exit 3", description="slab cast"),
                    _row((4, 5, 6), "2026-09-29T16:40:00+08:00", None, (1200, 1600),
                         caption="waterproofing")]),
        ("U3-12", [_row((7, 8, 9), "2026-09-30T10:58:00+08:00", "U3-12")]),
        ("RS", []),
    ]
    wb = load_workbook(io.BytesIO(xls.generate_photo_excel(sections, "30 Jul - 30 Sep 2026")))
    assert wb.sheetnames == ["Exit 3", "U3-12", "RS"]

    exit3 = wb["Exit 3"]
    assert exit3["A1"].value == "Exit 3 — 30 Jul - 30 Sep 2026"
    # Exactly Date, Time, Description, Photo — nothing else.
    assert [c.value for c in exit3[3]] == ["Date", "Time", "Description", "Photo"]
    assert exit3.max_column == 4
    assert exit3["B4"].value == "09:15" and exit3["C4"].value == "slab cast"
    assert exit3["B5"].value == "16:40" and exit3["C5"].value == "waterproofing"
    assert {img.anchor._from.col for img in exit3._images} == {3}      # column D
    assert sorted(img.anchor._from.row for img in exit3._images) == [3, 4]
    assert exit3.row_dimensions[5].height > exit3.row_dimensions[4].height   # portrait

    assert len(wb["U3-12"]._images) == 1
    assert wb["RS"]._images == []
    assert "No photos found" in wb["RS"]["A4"].value


@pytest.mark.parametrize("names, expected", [
    (["Exit 3"], ["Exit 3"]),
    (["U3/12: slab"], ["U3-12- slab"]),
    (["A" * 40], ["A" * 31]),
    (["Exit 3", "exit 3"], ["Exit 3", "exit 3 (2)"]),
])
def test_tab_names_are_valid_and_unique(names, expected):
    used = set()
    assert [xls._sheet_name(n, used) for n in names] == expected


# ── /ask routing ──────────────────────────────────────────────────────────────

class TestPhotoAsk:
    @pytest.fixture
    def out(self, monkeypatch):
        box = []

        async def send(group_id, text):
            box.append(("text", text))

        async def send_doc(group_id, data, filename, caption=""):
            box.append(("file", filename, caption))

        async def no_catch_up(_g):
            pass

        monkeypatch.setattr(cmd, "send_message", send)
        monkeypatch.setattr(cmd, "send_document", send_doc)
        monkeypatch.setattr(ib, "catch_up", no_catch_up)
        return box

    def test_several_locations_come_back_as_one_file_with_a_count_per_tab(self, out, monkeypatch):
        monkeypatch.setattr(photos, "find_photos", lambda g, q: (
            [("Exit 3", [{"id": 1}, {"id": 2}]), ("U3-12", [{"id": 3}]), ("RS", [])],
            "30 Jul - 30 Sep 2026", False))
        built = {}
        monkeypatch.setattr(cmd.xls, "generate_photo_excel",
                            lambda sections, period: built.update(s=sections, p=period) or b"xlsx")
        asyncio.run(cmd.handle_ask(SITE_GROUP, "pictures of Exit 3, U3-12 and RS last two months"))
        assert [loc for loc, _ in built["s"]] == ["Exit 3", "U3-12", "RS"]
        assert out[-1] == ("file", "Photos_Exit_3_U3_12_RS.xlsx",
                           "📷 3 photos, 30 Jul - 30 Sep 2026\n• Exit 3: 2\n• U3-12: 1\n• RS: 0")

    def test_one_location_keeps_a_one_line_caption(self, out, monkeypatch):
        monkeypatch.setattr(photos, "find_photos", lambda g, q: (
            [("Vent Shaft", [{"id": 1}])], "30 Sep 2026", False))
        monkeypatch.setattr(cmd.xls, "generate_photo_excel", lambda s, p: b"xlsx")
        asyncio.run(cmd.handle_ask(SITE_GROUP, "photo of Vent Shaft today"))
        assert out[-1] == ("file", "Photos_Vent_Shaft.xlsx", "📷 1 photo, 30 Sep 2026")

    def test_no_photos_anywhere_says_so_and_when_saving_started(self, out, monkeypatch):
        monkeypatch.setattr(photos, "find_photos", lambda g, q: ([("Exit 9", [])], "", False))
        asyncio.run(cmd.handle_ask(SITE_GROUP, "pictures of Exit 9"))
        assert "couldn't find any photos" in out[-1][1]
        assert "30 Sep 2026" in out[-1][1]

    def test_an_ordinary_question_does_not_take_the_photo_path(self, out, monkeypatch):
        import ai_handler
        monkeypatch.setattr(photos, "find_photos",
                            lambda *a: pytest.fail("not a photo question"))
        monkeypatch.setattr(ai_handler, "answer_query", lambda g, q: "answer")
        asyncio.run(cmd.handle_ask(SITE_GROUP, "when was U3-38 cast?"))
        assert out[-1] == ("text", "🤖 answer")
